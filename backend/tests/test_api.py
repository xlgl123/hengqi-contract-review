from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document

TEST_DIR = Path(tempfile.mkdtemp(prefix="contract-mvp-test-"))
os.environ["DATABASE_PATH"] = str(TEST_DIR / "test.db")
os.environ["AI_CREDENTIAL_PATH"] = str(TEST_DIR / "ai-settings.dpapi")
os.environ.pop("AI_API_KEY", None)
os.environ.pop("DEEPSEEK_API_KEY", None)

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app, credential_store  # noqa: E402


class APITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client = TestClient(app)

    def test_health_is_honest_about_ai_mode(self) -> None:
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["ai_configured"])

    def test_sample_contract_runs_complete_flow(self) -> None:
        sample = self.client.get("/api/sample").json()
        response = self.client.post(
            "/api/reviews",
            data={"text": sample["text"], "party_role": "service_provider"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertGreater(len(payload["document_blocks"]), 0)
        self.assertEqual(payload["engine_mode"], "rules_demo")
        self.assertGreater(payload["risk_bill"]["total_risks"], 0)
        self.assertGreater(payload["risk_bill"]["high"], 0)
        self.assertIn("未配置AI_API_KEY", payload["warnings"][0])

        stored = self.client.get(f"/api/reviews/{payload['review_id']}")
        self.assertEqual(stored.status_code, 200)
        self.assertEqual(stored.json()["review_id"], payload["review_id"])
        self.assertEqual(stored.json()["document_blocks"], [])
        with sqlite3.connect(TEST_DIR / "test.db") as connection:
            result_json = connection.execute(
                "SELECT result_json FROM reviews WHERE review_id = ?", (payload["review_id"],)
            ).fetchone()[0]
        self.assertNotIn('"document_blocks"', result_json)
        self.assertNotIn(payload["document_blocks"][0]["text"], result_json)

        recent = self.client.get("/api/reviews?limit=5")
        self.assertEqual(recent.status_code, 200)
        self.assertTrue(any(item["review_id"] == payload["review_id"] for item in recent.json()))
        summary = next(item for item in recent.json() if item["review_id"] == payload["review_id"])
        self.assertEqual(summary["total_risks"], payload["risk_bill"]["total_risks"])

        first_risk = payload["risks"][0]
        exported = self.client.post(
            f"/api/reviews/{payload['review_id']}/export/docx",
            json={"risk_statuses": {first_risk["risk_id"]: "adopted"}},
        )
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertEqual(
            exported.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertTrue(exported.content.startswith(b"PK"))
        self.assertIn(".docx", exported.headers["content-disposition"])
        report = Document(BytesIO(exported.content))
        report_text = "\n".join(
            paragraph.text for paragraph in report.paragraphs
        ) + "\n" + "\n".join(
            cell.text for table in report.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("合同风险审查报告", report_text)
        self.assertIn(first_risk["title"], report_text)
        self.assertIn("已采用建议", report_text)
        self.assertIn("法律依据", report_text)

    def test_export_rejects_invalid_status_and_missing_review(self) -> None:
        invalid = self.client.post(
            "/api/reviews/not-found/export/docx",
            json={"risk_statuses": {"risk-1": "finished"}},
        )
        self.assertEqual(invalid.status_code, 422)
        missing = self.client.post(
            "/api/reviews/not-found/export/docx",
            json={"risk_statuses": {}},
        )
        self.assertEqual(missing.status_code, 404)

    def test_short_text_is_rejected(self) -> None:
        response = self.client.post(
            "/api/reviews",
            data={"text": "太短", "party_role": "service_provider"},
        )
        self.assertEqual(response.status_code, 400)

    def test_legal_library_has_verified_official_sources(self) -> None:
        info = self.client.get("/api/legal-library")
        self.assertEqual(info.status_code, 200)
        payload = info.json()
        self.assertGreaterEqual(payload["basis_count"], 10)
        self.assertEqual(payload["verified_at"], "2026-07-16")

        bases = self.client.get("/api/legal-bases")
        self.assertEqual(bases.status_code, 200)
        self.assertEqual(len(bases.json()), payload["basis_count"])

    @patch("app.main.AIClient.test_connection", new_callable=AsyncMock)
    def test_ai_can_be_configured_and_cleared_from_local_ui(self, test_connection: AsyncMock) -> None:
        test_connection.return_value = {
            "model": "deepseek-ai/DeepSeek-V4-Flash",
            "total_tokens": 23,
        }
        self.client.delete("/api/settings/ai")
        configured = self.client.post(
            "/api/settings/ai",
            json={
                "provider": "siliconflow",
                "api_key": "test-key-not-a-real-secret-123",
                "model": "deepseek-ai/DeepSeek-V4-Flash",
                "remember": True,
            },
        )
        self.assertEqual(configured.status_code, 200, configured.text)
        payload = configured.json()
        self.assertTrue(payload["configured"])
        self.assertEqual(payload["provider"], "siliconflow")
        self.assertEqual(payload["storage"], "encrypted_local")
        self.assertEqual(payload["tested_model"], "deepseek-ai/DeepSeek-V4-Flash")
        self.assertNotIn("test-key-not-a-real-secret", configured.text)
        self.assertTrue(self.client.get("/api/health").json()["ai_configured"])
        self.assertTrue(credential_store.path.is_file())
        restored = credential_store.load()
        self.assertIsNotNone(restored)
        self.assertEqual(restored.api_key, "test-key-not-a-real-secret-123")
        restart_env = os.environ.copy()
        restart_env.pop("AI_API_KEY", None)
        restart_env.pop("DEEPSEEK_API_KEY", None)
        restart_env["AI_CREDENTIAL_PATH"] = str(credential_store.path)
        restart_env["DATABASE_PATH"] = str(TEST_DIR / "restart.db")
        restarted = subprocess.run(
            [
                sys.executable,
                "-c",
                (
                    "from app.main import ai_storage_mode, service;"
                    "print(ai_storage_mode);"
                    "print(service.settings.ai_api_key)"
                ),
            ],
            cwd=Path(__file__).resolve().parents[1],
            env=restart_env,
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertEqual(
            restarted.stdout.strip().splitlines(),
            ["encrypted_local", "test-key-not-a-real-secret-123"],
        )

        cleared = self.client.delete("/api/settings/ai")
        self.assertEqual(cleared.status_code, 200)
        self.assertFalse(cleared.json()["configured"])
        self.assertFalse(credential_store.path.exists())
        self.assertFalse(self.client.get("/api/health").json()["ai_configured"])


if __name__ == "__main__":
    unittest.main()
