from __future__ import annotations

import io
import unittest

from docx import Document

from app.config import get_settings
from app.parsers.service import parse_pasted_text, parse_upload


def build_minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    )
    return bytes(output)


class ParserTests(unittest.TestCase):
    def setUp(self) -> None:
        self.settings = get_settings()

    def test_pasted_text_has_stable_blocks(self) -> None:
        parsed = parse_pasted_text("技术服务合同\n第一条 服务内容\n乙方提供系统开发服务并交付源代码。")
        self.assertEqual(parsed.file_type, "pasted_text")
        self.assertEqual(parsed.blocks[0].block_id, "text-p1")
        self.assertGreaterEqual(len(parsed.blocks), 3)

    def test_docx_paragraphs_and_tables_keep_order(self) -> None:
        document = Document()
        document.add_paragraph("技术服务合同")
        document.add_paragraph("第一条 服务内容：乙方提供软件开发服务。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "合同金额"
        table.cell(0, 1).text = "100000元"
        stream = io.BytesIO()
        document.save(stream)

        parsed = parse_upload("test.docx", stream.getvalue(), self.settings)
        self.assertEqual(parsed.file_type, "docx")
        self.assertEqual([block.kind for block in parsed.blocks], ["paragraph", "paragraph", "table_row"])
        self.assertEqual(parsed.blocks[-1].block_id, "docx-t1-r1")

    def test_text_pdf_is_parsed_with_page_number(self) -> None:
        content = build_minimal_pdf("Technical Service Contract payment after acceptance without deadline.")
        parsed = parse_upload("test.pdf", content, self.settings)
        self.assertEqual(parsed.file_type, "pdf")
        self.assertEqual(parsed.page_count, 1)
        self.assertEqual(parsed.blocks[0].page, 1)


if __name__ == "__main__":
    unittest.main()

