from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas import ReviewResult, RiskStatus


STATUS_LABELS: dict[RiskStatus, str] = {
    "pending": "待处理",
    "adopted": "已采用建议",
    "communicated": "已沟通",
    "ignored": "暂时忽略",
}
SEVERITY_LABELS = {"high": "高风险", "medium": "中风险", "low": "低风险"}
LEGAL_NATURE_LABELS = {
    "compliance_risk": "法定合规风险",
    "enforceability_risk": "效力 / 执行风险",
    "agreement_gap": "重要约定缺失",
    "commercial_risk": "商业风险",
}
CATEGORY_LABELS = {
    "acceptance": "验收",
    "payment": "付款",
    "scope": "服务范围",
    "change_management": "需求变更",
    "intellectual_property": "知识产权",
    "liability": "违约责任",
    "termination": "合同解除",
    "dispute_resolution": "争议解决",
    "confidentiality": "保密",
    "data_compliance": "数据合规",
    "data_privacy": "数据与个人信息",
}

INK = "0B2545"
BLUE = "2E74B5"
MUTED = "667085"
LIGHT = "E8EEF5"
GRID = "CBD5E1"
RISK_COLORS = {"high": "B42318", "medium": "B54708", "low": "027A48"}


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_run_font(run, *, size: float = 11, color: str = INK, bold: bool = False) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, 14, 8, BLUE),
        "Heading 2": (13, 11, 6, BLUE),
        "Heading 3": (12, 8, 4, "1F4D78"),
    }
    for name, (size, before, after, color) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    _set_run_font(header.add_run("衡契 · 合同风险审查报告"), size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    _set_run_font(footer.add_run("第 "), size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    page_run = footer.add_run()
    _set_run_font(page_run, size=9, color=MUTED)
    page_run._r.extend([fld_char_begin, instr_text, fld_char_end])
    _set_run_font(footer.add_run(" 页"), size=9, color=MUTED)


def _add_label_paragraph(document: Document, label: str, value: str, *, color: str = INK) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    _set_run_font(paragraph.add_run(f"{label}："), bold=True, color=MUTED)
    _set_run_font(paragraph.add_run(value or "未提供"), color=color)


def _add_metadata_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        _set_cell_shading(cells[0], LIGHT)
        cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
        cells[1].paragraphs[0].paragraph_format.space_after = Pt(0)
        _set_run_font(cells[0].paragraphs[0].add_run(label), bold=True, color="1F4D78")
        _set_run_font(cells[1].paragraphs[0].add_run(value))
    _set_table_geometry(table, [1700, 7660])


def build_review_docx(result: ReviewResult, statuses: dict[str, RiskStatus]) -> bytes:
    """Generate a genuine Word OOXML report from a stored review result."""
    document = Document()
    _style_document(document)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    _set_run_font(title.add_run("合同风险审查报告"), size=24, color=INK, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    _set_run_font(subtitle.add_run(result.filename), size=12, color=MUTED)

    engine = (
        f"规则＋AI 联合审查 · {result.model_name or '未记录模型'}"
        if result.engine_mode == "hybrid_ai"
        else "AI 异常 · 已安全降级"
        if result.engine_mode == "hybrid_fallback"
        else "本地规则演示模式"
    )
    _add_metadata_table(
        document,
        [
            ("合同类型", result.summary.contract_type),
            ("我方身份", result.summary.our_role),
            ("相对方", result.summary.counterparty_role),
            ("审查引擎", engine),
            ("法律依据库", f"{result.legal_library_version}（核验日期：{result.legal_library_verified_at or '未记录'}）"),
        ],
    )

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    note.paragraph_format.space_after = Pt(12)
    _set_run_font(note.add_run("重要提示："), bold=True, color="B54708")
    _set_run_font(note.add_run("本报告用于签约前经营风险辅助识别，不替代律师针对具体交易出具的法律意见。"), color="7A5A00")

    document.add_heading("一、审查总览", level=1)
    _add_metadata_table(
        document,
        [
            ("风险合计", str(result.risk_bill.total_risks)),
            ("风险分布", f"高风险 {result.risk_bill.high} 项 · 中风险 {result.risk_bill.medium} 项 · 低风险 {result.risk_bill.low} 项"),
            ("付款风险", result.risk_bill.payment_exposure),
            ("验收风险", result.risk_bill.acceptance_exposure),
            ("责任风险", result.risk_bill.liability_exposure),
            ("知识产权风险", result.risk_bill.ip_exposure),
        ],
    )

    if result.summary.must_fix:
        document.add_heading("签约前优先处理", level=2)
        for index, item in enumerate(result.summary.must_fix, start=1):
            _add_label_paragraph(document, f"优先 {index}", item, color="B42318")

    document.add_heading("二、逐项风险与修改方案", level=1)
    if not result.risks:
        document.add_paragraph("本次审查未识别出风险项。")

    for index, risk in enumerate(result.risks, start=1):
        heading = document.add_heading(level=2)
        heading.paragraph_format.keep_with_next = True
        _set_run_font(
            heading.add_run(f"{index}. [{SEVERITY_LABELS[risk.severity]}] {risk.title}"),
            size=13,
            color=RISK_COLORS[risk.severity],
            bold=True,
        )
        status = statuses.get(risk.risk_id, "pending")
        _add_metadata_table(
            document,
            [
                ("处理状态", STATUS_LABELS[status]),
                ("风险领域", CATEGORY_LABELS.get(risk.category, risk.category)),
                ("法律性质", LEGAL_NATURE_LABELS[risk.legal_nature]),
            ],
        )
        _add_label_paragraph(document, "原文引用", risk.quote or "合同中未约定")
        _add_label_paragraph(document, "法律判断", risk.legal_conclusion)
        _add_label_paragraph(document, "风险说明", risk.legal_issue)
        _add_label_paragraph(document, "经营后果", risk.business_impact)

        document.add_heading("修改与谈判方案", level=3)
        _add_label_paragraph(document, "建议修改", risk.ideal_revision)
        _add_label_paragraph(document, "可接受方案", risk.compromise_revision)
        _add_label_paragraph(document, "最低底线", risk.bottom_line, color="B42318")
        _add_label_paragraph(document, "沟通话术", risk.negotiation_message)

        document.add_heading("法律依据", level=3)
        if risk.legal_bases:
            for basis_index, basis in enumerate(risk.legal_bases, start=1):
                _add_label_paragraph(
                    document,
                    f"依据 {basis_index}",
                    f"{basis.document_name}{basis.article_number}（{basis.effective_status}）",
                    color="1F4D78",
                )
                _add_label_paragraph(document, "条文摘要", basis.article_excerpt)
                _add_label_paragraph(document, "适用说明", basis.application_note)
                if basis.limitations:
                    _add_label_paragraph(document, "适用边界", "；".join(basis.limitations))
                _add_label_paragraph(document, "官方来源", basis.official_url, color=BLUE)
        else:
            document.add_paragraph("暂无直接法条依据；该项属于经营风险判断。")

        if index < len(result.risks):
            separator = document.add_paragraph()
            separator.paragraph_format.space_before = Pt(5)
            separator.paragraph_format.space_after = Pt(0)
            separator_format = separator._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "6")
            bottom.set(qn("w:color"), GRID)
            borders.append(bottom)
            separator_format.append(borders)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
